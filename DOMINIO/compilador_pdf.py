# DOMINIO/compilador_pdf.py
import os
import subprocess
import openpyxl
from PyPDF2 import PdfMerger


class CompiladorPDF:
    @staticmethod
    def recalcular_y_exportar_pdf(
        ruta_xlsx_modificado: str, ruta_pdf_destino: str
    ) -> str:

        wb = openpyxl.load_workbook(ruta_xlsx_modificado)
        hojas_a_conservar = [
            nombre
            for nombre in wb.sheetnames
            if not nombre.startswith("(") and "HOJA DE" not in nombre.upper()
        ]

        for sheet in wb.sheetnames:
            if sheet not in hojas_a_conservar and len(wb.sheetnames) > 1:
                del wb[sheet]

        ruta_temp_xlsx = ruta_xlsx_modificado.replace(".xlsx", "_clean.xlsx")
        wb.save(ruta_temp_xlsx)
        wb.close()

        carpeta_salida = os.path.dirname(ruta_pdf_destino)
        cmd = [
            "libreoffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            carpeta_salida,
            ruta_temp_xlsx,
        ]

        try:
            subprocess.run(
                cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE
            )
        except (subprocess.CalledProcessError, FileNotFoundError):
            pass

        if os.path.exists(ruta_temp_xlsx):
            os.remove(ruta_temp_xlsx)

        return ruta_pdf_destino

    @staticmethod
    def consolidar_expediente(lista_pdfs: list, ruta_pdf_final: str) -> str:

        merger = PdfMerger()
        for pdf in lista_pdfs:
            if os.path.exists(pdf):
                merger.append(pdf)

        os.makedirs(os.path.dirname(ruta_pdf_final), exist_ok=True)
        merger.write(ruta_pdf_final)
        merger.close()

        return ruta_pdf_final
